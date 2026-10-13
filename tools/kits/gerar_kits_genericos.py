"""Gera um Kit de Arranque para cada especificação em tools/kits/specs/*.json.

Cada kit vira uma planilha .xlsx com: Como usar, Painel, Diagnóstico, Plano de ação, Procedimento e
as abas de controle próprias do problema (listas suspensas, formatos, linhas de exemplo). As fórmulas
são criadas aqui, nunca pela especificação, então todos os kits funcionam igual no Excel, Google
Sheets e LibreOffice.

Também grava:
- apps/api/kits/indice.json: slug → título e subdivisões (lido por apps/api/src/constants/kits.js);
- apps/web/src/data/kitSubdivisions.json: subdivisões com kit (lido por apps/web/src/data/hub.js).

Uso: python tools/kits/gerar_kits_genericos.py   (valida tudo antes de gravar; falha se houver erro)
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from gerar_kits import (AMARELO, AMARELO_CLARO, DATA, INDIGO, MOEDA, OUT, VERDE, VERDE_CLARO,  # noqa: E402
                        VERMELHO, VERMELHO_CLARO, cabecalho, como_usar, destaque, formatos, lista, painel, titulo)

RAIZ = Path(__file__).resolve().parents[2]
SPECS = Path(__file__).resolve().parent / "specs"
WEB_LISTA = RAIZ / "apps" / "web" / "src" / "data" / "kitSubdivisions.json"
HUB = RAIZ / "apps" / "web" / "src" / "data" / "hub.js"
# subdivisões que já têm kit feito à mão em gerar_kits.py (têm prioridade)
MANUAIS = {"controle-de-estoque": ["Controle de estoque", "Gestão de estoque"], "gestao-de-pedidos": ["Gestão de pedidos"],
           "controle-de-vencimentos": ["Controle de vencimentos"]}
RESERVADAS = {"como usar", "painel", "diagnóstico", "plano de ação", "procedimento"}
FORMATO = {"texto": None, "numero": "#,##0.##", "moeda": MOEDA, "data": DATA, "lista": None}
INICIO = 5  # primeira linha de dados nas abas


def subdivisoes_do_site() -> set[str]:
    """Os itens de cada área em hub.js (os blocos `items: [ ... ]`)."""
    texto = HUB.read_text(encoding="utf-8")
    return {item for bloco in re.findall(r"items:\s*\[(.*?)\]", texto, re.S) for item in re.findall(r"'([^']+)'", bloco)}


def validar(spec: dict, arquivo: str) -> list[str]:
    erros = []
    for campo in ("slug", "titulo", "subdivisoes", "resumo", "como_usar", "diagnostico", "plano", "procedimento", "abas", "dicas"):
        if not spec.get(campo):
            erros.append(f"{arquivo}: falta '{campo}'")
    if erros:
        return erros
    if spec["slug"] != Path(arquivo).stem:
        erros.append(f"{arquivo}: slug '{spec['slug']}' diferente do nome do arquivo")
    for p in spec["plano"]:
        if not isinstance(p.get("prazo_dias"), (int, float)) or not p.get("acao"):
            erros.append(f"{arquivo}: item do plano sem 'acao' ou 'prazo_dias' numérico: {p}")
    nomes = set()
    for aba in spec["abas"]:
        nome = str(aba.get("nome", ""))
        if not nome or len(nome) > 28 or re.search(r"[/\\?*\[\]:]", nome) or nome.lower() in RESERVADAS or nome in nomes:
            erros.append(f"{arquivo}: nome de aba inválido ou repetido: '{nome}'")
        nomes.add(nome)
        colunas = aba.get("colunas") or []
        if not 1 <= len(colunas) <= 10:
            erros.append(f"{arquivo}: aba '{nome}' com {len(colunas)} colunas")
        for c in colunas:
            if c.get("tipo") not in FORMATO:
                erros.append(f"{arquivo}: coluna '{c.get('nome')}' com tipo '{c.get('tipo')}'")
            if c.get("tipo") == "lista" and not c.get("opcoes"):
                erros.append(f"{arquivo}: coluna lista '{c.get('nome')}' sem opcoes")
        for linha in aba.get("exemplos") or []:
            if len(linha) != len(colunas):
                erros.append(f"{arquivo}: aba '{nome}': exemplo com {len(linha)} valores para {len(colunas)} colunas")
                continue
            for valor, c in zip(linha, colunas):
                t = c.get("tipo")
                if t in ("numero", "moeda") and not isinstance(valor, (int, float)):
                    erros.append(f"{arquivo}: '{c['nome']}' espera número, veio {valor!r}")
                if t == "data" and not re.fullmatch(r"[+-]?\d+", str(valor)):
                    erros.append(f"{arquivo}: '{c['nome']}' espera data relativa (+5/-3/0), veio {valor!r}")
                if t == "lista" and valor not in (c.get("opcoes") or []):
                    erros.append(f"{arquivo}: '{c['nome']}' valor {valor!r} fora das opções")
    return erros


def _q(nome: str) -> str:
    return "'" + nome.replace("'", "''") + "'"


def diagnostico(wb, perguntas):
    ws = wb.create_sheet("Diagnóstico")
    titulo(ws, "Diagnóstico rápido", "Responda com sinceridade: mostra onde começar.")
    cols = [("Pergunta", 80, None, False), ("Resposta", 16, None, True), ("Observação", 40, None, True)]
    cabecalho(ws, 4, cols)
    for i, p in enumerate(perguntas):
        c = ws.cell(row=INICIO + i, column=1, value=p)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    fim = INICIO + len(perguntas) - 1
    lista(ws, f"B{INICIO}:B{fim}", ["Sim", "Em parte", "Não"])
    destaque(ws, f"B{INICIO}:B{fim}", f'B{INICIO}="Não"', VERMELHO, VERMELHO_CLARO)
    destaque(ws, f"B{INICIO}:B{fim}", f'B{INICIO}="Em parte"', AMARELO, AMARELO_CLARO)
    destaque(ws, f"B{INICIO}:B{fim}", f'B{INICIO}="Sim"', VERDE, VERDE_CLARO)
    return ws, fim


def plano_de_acao(wb, plano):
    ws = wb.create_sheet("Plano de ação")
    titulo(ws, "Plano de ação", "Prazos contados a partir de hoje. Ajuste as datas se precisar.")
    cols = [("Etapa", 14, None, True), ("Ação", 60, None, True), ("Responsável", 18, None, True),
            ("Prazo", 13, DATA, True), ("Status", 15, None, True), ("Situação", 16, None, False)]
    cabecalho(ws, 4, cols)
    formatos(ws, INICIO, cols)
    for i, p in enumerate(plano):
        r = INICIO + i
        ws.cell(row=r, column=1, value=p.get("etapa", ""))
        ws.cell(row=r, column=2, value=p["acao"]).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=3, value=p.get("responsavel", ""))
        ws.cell(row=r, column=4, value=f"=TODAY()+{int(p['prazo_dias'])}")
        ws.cell(row=r, column=5, value="A fazer")
    fim = INICIO + 199
    for r in range(INICIO, fim + 1):
        ws.cell(row=r, column=6, value=f'=IF(B{r}="","",IF(E{r}="Feito","Concluída",IF(D{r}<TODAY(),"Atrasada","No prazo")))')
    lista(ws, f"E{INICIO}:E{fim}", ["A fazer", "Em andamento", "Feito"])
    destaque(ws, f"F{INICIO}:F{fim}", f'F{INICIO}="Atrasada"', VERMELHO, VERMELHO_CLARO)
    destaque(ws, f"F{INICIO}:F{fim}", f'F{INICIO}="Concluída"', VERDE, VERDE_CLARO)
    return ws, fim


def procedimento(wb, passos):
    ws = wb.create_sheet("Procedimento")
    titulo(ws, "Procedimento passo a passo", "Use como padrão da empresa (POP). Imprima ou compartilhe com a equipe.")
    cols = [("Nº", 6, None, False), ("Passo", 28, None, True), ("Como fazer", 60, None, True),
            ("Cuidado", 45, None, True), ("Feito?", 10, None, True)]
    cabecalho(ws, 4, cols)
    for i, p in enumerate(passos):
        r = INICIO + i
        ws.cell(row=r, column=1, value=i + 1).font = Font(bold=True, color=INDIGO)
        for col, chave in ((2, "passo"), (3, "como"), (4, "cuidado")):
            ws.cell(row=r, column=col, value=p.get(chave, "")).alignment = Alignment(wrap_text=True, vertical="top")
    fim = INICIO + len(passos) - 1
    lista(ws, f"E{INICIO}:E{fim}", ["Sim", "Não"])
    destaque(ws, f"E{INICIO}:E{fim}", f'E{INICIO}="Sim"', VERDE, VERDE_CLARO)
    return ws, fim


def aba_controle(wb, aba):
    ws = wb.create_sheet(aba["nome"])
    titulo(ws, aba["nome"], aba.get("descricao", ""))
    cols = []
    for c in aba["colunas"]:
        largura = max(14, min(40, len(c["nome"]) + 6))
        cols.append((c["nome"], largura, FORMATO[c["tipo"]], True))
    cabecalho(ws, 4, cols)
    formatos(ws, INICIO, cols)
    for i, linha in enumerate(aba.get("exemplos") or []):
        for j, (valor, c) in enumerate(zip(linha, aba["colunas"]), start=1):
            if c["tipo"] == "data":
                valor = f"=TODAY()+({int(valor)})"
            ws.cell(row=INICIO + i, column=j, value=valor)
    fim = INICIO + 499
    for j, c in enumerate(aba["colunas"], start=1):
        if c["tipo"] == "lista":
            lista(ws, f"{get_column_letter(j)}{INICIO}:{get_column_letter(j)}{fim}", c["opcoes"])
    return ws, fim


def gerar(spec: dict) -> Path:
    wb = Workbook()
    como_usar(wb, spec["titulo"], [spec["resumo"], *spec["como_usar"]], spec["dicas"])
    _, fim_d = diagnostico(wb, spec["diagnostico"])
    _, fim_p = plano_de_acao(wb, spec["plano"])
    _, fim_s = procedimento(wb, spec["procedimento"])
    indicadores = [
        ("Ações no plano", f"=COUNTA({_q('Plano de ação')}!B{INICIO}:B{fim_p})", None),
        ("Ações concluídas", f'=COUNTIF({_q("Plano de ação")}!F{INICIO}:F{fim_p},"Concluída")', None),
        ("Ações atrasadas", f'=COUNTIF({_q("Plano de ação")}!F{INICIO}:F{fim_p},"Atrasada")', None),
        ("Plano concluído", f'=IF(COUNTA({_q("Plano de ação")}!B{INICIO}:B{fim_p})=0,0,COUNTIF({_q("Plano de ação")}!F{INICIO}:F{fim_p},"Concluída")/COUNTA({_q("Plano de ação")}!B{INICIO}:B{fim_p}))', "0%"),
        ("Diagnóstico: respostas \"Não\"", f'=COUNTIF({_q("Diagnóstico")}!B{INICIO}:B{fim_d},"Não")', None),
        ("Procedimento: passos feitos", f'=COUNTIF({_q("Procedimento")}!E{INICIO}:E{fim_s},"Sim")', None),
    ]
    for aba in spec["abas"]:
        _, fim = aba_controle(wb, aba)
        nome = _q(aba["nome"])
        indicadores.append((f"{aba['nome']}: registros", f"=COUNTA({nome}!A{INICIO}:A{fim})", None))
        for j, c in enumerate(aba["colunas"], start=1):
            col = get_column_letter(j)
            if c["tipo"] == "moeda":
                indicadores.append((f"{aba['nome']}: total de {c['nome']}", f"=SUM({nome}!{col}{INICIO}:{col}{fim})", MOEDA))
        lista_col = next(((j, c) for j, c in enumerate(aba["colunas"], start=1) if c["tipo"] == "lista"), None)
        if lista_col:
            j, c = lista_col
            col = get_column_letter(j)
            for op in c["opcoes"][:5]:
                indicadores.append((f"{aba['nome']}: {c['nome']} = {op}", f'=COUNTIF({nome}!{col}{INICIO}:{col}{fim},"{op}")', None))
    painel(wb, f"Painel · {spec['titulo']}", indicadores)
    destino = OUT / f"{spec['slug']}.xlsx"
    wb.save(destino)
    return destino


def main() -> int:
    specs, erros = [], []
    for arquivo in sorted(SPECS.glob("*.json")):
        try:
            spec = json.loads(arquivo.read_text(encoding="utf-8"))
        except json.JSONDecodeError as e:
            erros.append(f"{arquivo.name}: JSON inválido ({e})")
            continue
        erros += validar(spec, arquivo.name)
        specs.append(spec)
    site = subdivisoes_do_site()
    cobertas = {s for subs in MANUAIS.values() for s in subs}
    for spec in specs:
        for s in spec.get("subdivisoes", []):
            if s not in site:
                erros.append(f"{spec.get('slug')}: subdivisão '{s}' não existe em hub.js")
            if s in cobertas:
                erros.append(f"{spec.get('slug')}: subdivisão '{s}' já tem kit")
            cobertas.add(s)
    if erros:
        print("\n".join(erros))
        print(f"{len(erros)} erro(s): nada foi gravado.")
        return 1

    indice = {}
    for spec in specs:
        gerar(spec)
        indice[spec["slug"]] = {"title": f"Kit de Arranque · {spec['titulo']}", "subdivisions": spec["subdivisoes"]}
    (OUT / "indice.json").write_text(json.dumps(indice, ensure_ascii=False, indent=1), encoding="utf-8")
    WEB_LISTA.write_text(json.dumps(sorted(cobertas), ensure_ascii=False, indent=1), encoding="utf-8")
    faltando = sorted(site - cobertas)
    print(f"{len(specs)} kits gerados · {len(cobertas)} de {len(site)} subdivisões com kit")
    if faltando:
        print("Sem kit:", ", ".join(faltando))
    return 0


if __name__ == "__main__":
    sys.exit(main())
