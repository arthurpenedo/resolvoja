"""Gera as planilhas do Kit de Arranque entregues logo após o pagamento.

Cada kit é uma planilha .xlsx pronta para uso (abre no Excel, Google Sheets
e LibreOffice): aba "Como usar", abas de cadastro com listas suspensas,
fórmulas já preenchidas, formatação condicional e um painel de resumo.

Só funções compatíveis com as três ferramentas (SOMASES, SE, PROCV etc.,
gravadas com o nome em inglês, como o formato .xlsx exige).

Uso: python tools/kits/gerar_kits.py   (grava em apps/api/kits/)
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[2] / "apps" / "api" / "kits"
LINHAS = 500  # linhas preparadas com fórmula em cada aba de cadastro

INDIGO = "4F46E5"
INDIGO_CLARO = "EEF2FF"
CINZA = "6B7280"
VERMELHO, VERMELHO_CLARO = "B91C1C", "FEE2E2"
AMARELO, AMARELO_CLARO = "92400E", "FEF3C7"
VERDE, VERDE_CLARO = "166534", "DCFCE7"

BORDA = Border(bottom=Side(style="thin", color="E5E7EB"))
MOEDA = '"R$" #,##0.00'
DATA = "DD/MM/YYYY"
MOV = "'Movimentações'"  # nome com acento: sempre entre aspas nas fórmulas


def rel(dias):
    """Data de exemplo relativa a hoje, para a demonstração nunca ficar velha."""
    return f"=TODAY()+({dias})"


def titulo(ws, texto, subtitulo):
    ws["A1"] = texto
    ws["A1"].font = Font(size=18, bold=True, color=INDIGO)
    ws["A2"] = subtitulo
    ws["A2"].font = Font(size=10, italic=True, color=CINZA)
    ws.row_dimensions[1].height = 28


def cabecalho(ws, linha, colunas):
    """colunas: lista de (nome, largura, formato_ou_None, editavel)."""
    for i, (nome, largura, _fmt, editavel) in enumerate(colunas, start=1):
        c = ws.cell(row=linha, column=i, value=nome)
        c.font = Font(bold=True, color="FFFFFF" if editavel else INDIGO)
        c.fill = PatternFill("solid", fgColor=INDIGO if editavel else INDIGO_CLARO)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.row_dimensions[linha].height = 32
    ws.freeze_panes = ws.cell(row=linha + 1, column=1)


def formatos(ws, linha_ini, colunas):
    for i, (_n, _l, fmt, _e) in enumerate(colunas, start=1):
        for r in range(linha_ini, linha_ini + LINHAS):
            c = ws.cell(row=r, column=i)
            if fmt:
                c.number_format = fmt
            c.border = BORDA


def lista(ws, intervalo, opcoes=None, formula=None):
    dv = DataValidation(type="list", formula1=formula or '"' + ",".join(opcoes) + '"', allow_blank=True)
    dv.error = "Escolha um valor da lista."
    dv.errorTitle = "Valor inválido"
    ws.add_data_validation(dv)
    dv.add(intervalo)


def destaque(ws, intervalo, formula, cor_texto, cor_fundo):
    ws.conditional_formatting.add(
        intervalo,
        FormulaRule(formula=[formula], font=Font(bold=True, color=cor_texto),
                    fill=PatternFill("solid", fgColor=cor_fundo)),
    )


def como_usar(wb, nome_kit, passos, dicas):
    ws = wb.active
    ws.title = "Como usar"
    titulo(ws, f"Kit de Arranque · {nome_kit}", "Resolvo Já · resolvoja.com")
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 100
    ws["B4"] = "Passo a passo"
    ws["B4"].font = Font(size=13, bold=True)
    r = 5
    for i, p in enumerate(passos, start=1):
        ws.cell(row=r, column=1, value=i).font = Font(bold=True, color=INDIGO)
        ws.cell(row=r, column=2, value=p).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    r += 1
    ws.cell(row=r, column=2, value="Legenda das colunas").font = Font(size=13, bold=True)
    r += 1
    c = ws.cell(row=r, column=2, value="Cabeçalho azul escuro: você preenche.")
    c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=INDIGO)
    r += 1
    c = ws.cell(row=r, column=2, value="Cabeçalho azul claro: calculado automaticamente, não precisa mexer.")
    c.font, c.fill = Font(bold=True, color=INDIGO), PatternFill("solid", fgColor=INDIGO_CLARO)
    r += 2
    ws.cell(row=r, column=2, value="Dicas").font = Font(size=13, bold=True)
    r += 1
    for d in dicas:
        ws.cell(row=r, column=2, value=f"• {d}").alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    r += 1
    ws.cell(row=r, column=2, value=(
        "Os dados que já vêm na planilha são exemplos para você ver tudo funcionando. "
        "Apague o conteúdo das linhas de exemplo (só as colunas de cabeçalho escuro) e comece a usar."
    )).font = Font(italic=True, color=CINZA)
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True)
    ws.sheet_view.showGridLines = False
    return ws


def painel(wb, titulo_painel, indicadores, linha_ini=4):
    """indicadores: lista de (rótulo, fórmula, formato)."""
    ws = wb.create_sheet("Painel", 1)
    titulo(ws, titulo_painel, "Atualiza sozinho conforme você preenche as outras abas.")
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 22
    for i, (rotulo, formula, fmt) in enumerate(indicadores):
        r = linha_ini + i
        a = ws.cell(row=r, column=1, value=rotulo)
        b = ws.cell(row=r, column=2, value=formula)
        a.font = Font(bold=True)
        b.font = Font(size=13, bold=True, color=INDIGO)
        b.alignment = Alignment(horizontal="right")
        if fmt:
            b.number_format = fmt
        a.border = b.border = BORDA
        ws.row_dimensions[r].height = 24
    ws.sheet_view.showGridLines = False
    return ws


# --------------------------------------------------------------------------
# Kit 1: Controle de estoque
# --------------------------------------------------------------------------
def kit_estoque():
    wb = Workbook()
    como_usar(wb, "Controle de Estoque", [
        "Na aba Produtos, cadastre cada item com um código único, o estoque que você tem hoje (Estoque inicial) "
        "e o Estoque mínimo, a quantidade abaixo da qual você precisa comprar de novo.",
        "Toda vez que um produto entrar (compra, devolução) ou sair (venda, perda, uso interno), registre uma linha "
        "na aba Movimentações. Escolha o código na lista; o nome do produto aparece sozinho.",
        "Na aba Produtos, o Saldo atual e a Situação são recalculados na hora: verde = OK, amarelo = repor, "
        "vermelho = sem estoque.",
        "Use a coluna Sugestão de compra para montar o pedido ao fornecedor: ela indica quanto comprar para "
        "voltar ao dobro do estoque mínimo.",
        "O Painel resume quantos itens precisam de reposição e quanto dinheiro está parado em estoque.",
    ], [
        "Faça um inventário (contagem física) uma vez por mês. Se a contagem não bater com o Saldo atual, "
        "registre a diferença como Entrada ou Saída com a observação \"ajuste de inventário\".",
        "Para definir o estoque mínimo: média vendida por dia × dias que o fornecedor leva para entregar, com uma folga.",
        "Itens com muito valor parado e pouca saída são candidatos a promoção ou a comprar menos na próxima vez.",
    ])

    p = wb.create_sheet("Produtos")
    titulo(p, "Produtos", "Um produto por linha. O código precisa ser único.")
    cols = [
        ("Código", 10, None, True), ("Produto", 30, None, True), ("Categoria", 16, None, True),
        ("Unidade", 10, None, True), ("Custo unitário", 14, MOEDA, True), ("Preço de venda", 14, MOEDA, True),
        ("Estoque inicial", 12, "0", True), ("Estoque mínimo", 12, "0", True),
        ("Entradas", 11, "0", False), ("Saídas", 11, "0", False), ("Saldo atual", 11, "0", False),
        ("Situação", 15, None, False), ("Valor em estoque", 16, MOEDA, False), ("Sugestão de compra", 14, "0", False),
        ("Margem", 10, "0%", False),
    ]
    H = 4
    cabecalho(p, H, cols)
    formatos(p, H + 1, cols)
    lista(p, f"D{H + 1}:D{H + LINHAS}", ["un", "kg", "g", "L", "ml", "cx", "pct", "m", "par"])
    for r in range(H + 1, H + 1 + LINHAS):
        mov = f"{MOV}!$B$5:$B$5000"
        p[f"I{r}"] = f'=IF(A{r}="","",SUMIFS({MOV}!$E$5:$E$5000,{mov},A{r},{MOV}!$D$5:$D$5000,"Entrada"))'
        p[f"J{r}"] = f'=IF(A{r}="","",SUMIFS({MOV}!$E$5:$E$5000,{mov},A{r},{MOV}!$D$5:$D$5000,"Saída"))'
        p[f"K{r}"] = f'=IF(A{r}="","",G{r}+I{r}-J{r})'
        p[f"L{r}"] = f'=IF(A{r}="","",IF(K{r}<=0,"SEM ESTOQUE",IF(K{r}<=H{r},"REPOR","OK")))'
        p[f"M{r}"] = f'=IF(A{r}="","",MAX(K{r},0)*E{r})'
        p[f"N{r}"] = f'=IF(A{r}="","",IF(K{r}<=H{r},MAX(2*H{r}-K{r},0),0))'
        p[f"O{r}"] = f'=IF(OR(A{r}="",F{r}=0,F{r}=""),"",(F{r}-E{r})/F{r})'
    faixa = f"L{H + 1}:L{H + LINHAS}"
    destaque(p, faixa, f'L{H + 1}="SEM ESTOQUE"', VERMELHO, VERMELHO_CLARO)
    destaque(p, faixa, f'L{H + 1}="REPOR"', AMARELO, AMARELO_CLARO)
    destaque(p, faixa, f'L{H + 1}="OK"', VERDE, VERDE_CLARO)
    exemplos = [
        ("P001", "Café torrado 500 g", "Mercearia", "pct", 14.9, 24.9, 40, 15),
        ("P002", "Açúcar refinado 1 kg", "Mercearia", "pct", 4.2, 6.99, 25, 20),
        ("P003", "Leite integral 1 L", "Laticínios", "L", 3.8, 5.99, 60, 30),
        ("P004", "Copo descartável 200 ml", "Descartáveis", "pct", 6.5, 11.0, 12, 10),
        ("P005", "Guardanapo", "Descartáveis", "pct", 2.1, 4.5, 30, 10),
        ("P006", "Chocolate em pó 400 g", "Mercearia", "un", 9.8, 16.9, 8, 6),
        ("P007", "Filtro de papel 103", "Mercearia", "cx", 3.9, 7.5, 18, 8),
        ("P008", "Pão de queijo congelado 1 kg", "Congelados", "kg", 18.0, 32.0, 10, 5),
    ]
    for i, linha in enumerate(exemplos):
        for j, v in enumerate(linha, start=1):
            p.cell(row=H + 1 + i, column=j, value=v)

    m = wb.create_sheet("Movimentações")
    titulo(m, "Movimentações", "Uma linha para cada entrada ou saída de produto.")
    mcols = [
        ("Data", 12, DATA, True), ("Código", 10, None, True), ("Produto", 30, None, False),
        ("Tipo", 10, None, True), ("Quantidade", 11, "0", True), ("Observação", 34, None, True),
    ]
    cabecalho(m, H, mcols)
    formatos(m, H + 1, mcols)
    lista(m, f"B{H + 1}:B{H + LINHAS}", formula=f"=Produtos!$A${H + 1}:$A${H + LINHAS}")
    lista(m, f"D{H + 1}:D{H + LINHAS}", ["Entrada", "Saída"])
    for r in range(H + 1, H + 1 + LINHAS):
        m[f"C{r}"] = f'=IF(B{r}="","",IFERROR(VLOOKUP(B{r},Produtos!$A$5:$B$504,2,FALSE),"código não cadastrado"))'
    movs = [
        (-20, "P001", "Saída", 12, "vendas da semana"), (-18, "P003", "Saída", 22, "vendas da semana"),
        (-15, "P002", "Saída", 9, ""), (-14, "P004", "Saída", 7, "evento"),
        (-12, "P001", "Entrada", 10, "compra fornecedor A"), (-10, "P006", "Saída", 5, ""),
        (-9, "P008", "Saída", 6, ""), (-7, "P003", "Saída", 18, ""),
        (-6, "P005", "Saída", 8, ""), (-5, "P007", "Saída", 4, ""),
        (-3, "P002", "Entrada", 12, "compra fornecedor B"), (-2, "P004", "Saída", 5, ""),
        (-1, "P008", "Saída", 4, ""), (-1, "P001", "Saída", 14, ""),
    ]
    for i, (dias, cod, tipo, qtd, obs) in enumerate(movs):
        r = H + 1 + i
        m[f"A{r}"] = rel(dias)
        m[f"A{r}"].number_format = DATA
        m[f"B{r}"], m[f"D{r}"], m[f"E{r}"], m[f"F{r}"] = cod, tipo, qtd, obs

    faixa_cod = f"Produtos!$A${H + 1}:$A${H + LINHAS}"
    faixa_sit = f"Produtos!$L${H + 1}:$L${H + LINHAS}"
    painel(wb, "Painel do Estoque", [
        ("Produtos cadastrados", f'=COUNTA({faixa_cod})', "0"),
        ("Itens para repor", f'=COUNTIF({faixa_sit},"REPOR")', "0"),
        ("Itens sem estoque", f'=COUNTIF({faixa_sit},"SEM ESTOQUE")', "0"),
        ("Valor total em estoque (custo)", f"=SUM(Produtos!$M${H + 1}:$M${H + LINHAS})", MOEDA),
        ("Unidades a comprar (sugestão)", f"=SUM(Produtos!$N${H + 1}:$N${H + LINHAS})", "0"),
        ("Movimentações nos últimos 30 dias",
         f'=COUNTIFS({MOV}!$A$5:$A$5000,">="&(TODAY()-30))', "0"),
    ])
    return wb


# --------------------------------------------------------------------------
# Kit 2: Gestão de pedidos
# --------------------------------------------------------------------------
def kit_pedidos():
    wb = Workbook()
    como_usar(wb, "Gestão de Pedidos", [
        "Na aba Catálogo, cadastre seus produtos ou serviços com o preço de venda.",
        "Na aba Pedidos, registre cada pedido. Se o pedido tiver mais de um produto, use uma linha por produto "
        "repetindo o mesmo Nº do pedido.",
        "Escolha o produto na lista: o preço aparece sozinho e o total é calculado (quantidade × preço − desconto).",
        "Atualize a coluna Status conforme o pedido anda: Novo → Em produção → Pronto → Entregue. "
        "Pedidos com entrega vencida e ainda não entregues ficam marcados em vermelho como ATRASADO.",
        "No Painel, digite o mês e o ano que quer analisar e veja faturamento, valores a receber e os produtos mais vendidos.",
    ], [
        "Marque Pago = Sim só quando o dinheiro entrar. O Painel mostra quanto ainda falta receber.",
        "Pedidos cancelados não entram no faturamento, mas continuam registrados para você entender os motivos.",
        "Reveja todo dia de manhã os pedidos com entrega para hoje e para amanhã (dá para filtrar a coluna Entrega).",
    ])

    H = 4
    c = wb.create_sheet("Catálogo")
    titulo(c, "Catálogo", "Seus produtos ou serviços e o preço de venda.")
    ccols = [("Produto / serviço", 34, None, True), ("Preço de venda", 14, MOEDA, True), ("Categoria", 18, None, True)]
    cabecalho(c, H, ccols)
    formatos(c, H + 1, ccols)
    catalogo = [
        ("Bolo de pote", 12.0, "Doces"), ("Brigadeiro gourmet (cento)", 150.0, "Doces"),
        ("Torta de limão inteira", 85.0, "Tortas"), ("Kit festa 20 pessoas", 260.0, "Kits"),
        ("Taxa de entrega", 10.0, "Serviços"),
    ]
    for i, linha in enumerate(catalogo):
        for j, v in enumerate(linha, start=1):
            c.cell(row=H + 1 + i, column=j, value=v)

    p = wb.create_sheet("Pedidos", 1)
    titulo(p, "Pedidos", "Uma linha por produto do pedido. Mesmo Nº = mesmo pedido.")
    pcols = [
        ("Nº pedido", 9, "0", True), ("Data do pedido", 12, DATA, True), ("Cliente", 22, None, True),
        ("Telefone", 15, None, True), ("Produto", 28, None, True), ("Qtd", 7, "0", True),
        ("Preço unitário", 13, MOEDA, False), ("Desconto (R$)", 12, MOEDA, True), ("Total", 13, MOEDA, False),
        ("Entrega", 12, DATA, True), ("Pagamento", 13, None, True), ("Pago?", 8, None, True),
        ("Status", 14, None, True), ("Alerta", 12, None, False),
    ]
    cabecalho(p, H, pcols)
    formatos(p, H + 1, pcols)
    lista(p, f"E{H + 1}:E{H + LINHAS}", formula=f"='Catálogo'!$A${H + 1}:$A${H + LINHAS}")
    lista(p, f"K{H + 1}:K{H + LINHAS}", ["Pix", "Dinheiro", "Cartão de crédito", "Cartão de débito", "Boleto", "Fiado"])
    lista(p, f"L{H + 1}:L{H + LINHAS}", ["Sim", "Não"])
    lista(p, f"M{H + 1}:M{H + LINHAS}", ["Novo", "Em produção", "Pronto", "Entregue", "Cancelado"])
    for r in range(H + 1, H + 1 + LINHAS):
        p[f"G{r}"] = f"=IF(E{r}=\"\",\"\",IFERROR(VLOOKUP(E{r},'Catálogo'!$A$5:$B$504,2,FALSE),0))"
        p[f"I{r}"] = f'=IF(E{r}="","",MAX(F{r}*G{r}-N(H{r}),0))'
        p[f"N{r}"] = (f'=IF(OR(E{r}="",M{r}="Entregue",M{r}="Cancelado",J{r}=""),"",'
                      f'IF(J{r}<TODAY(),"ATRASADO",IF(J{r}=TODAY(),"ENTREGA HOJE","")))')
    destaque(p, f"N{H + 1}:N{H + LINHAS}", f'N{H + 1}="ATRASADO"', VERMELHO, VERMELHO_CLARO)
    destaque(p, f"N{H + 1}:N{H + LINHAS}", f'N{H + 1}="ENTREGA HOJE"', AMARELO, AMARELO_CLARO)
    destaque(p, f"M{H + 1}:M{H + LINHAS}", f'M{H + 1}="Entregue"', VERDE, VERDE_CLARO)
    destaque(p, f"L{H + 1}:L{H + LINHAS}", f'L{H + 1}="Não"', AMARELO, AMARELO_CLARO)
    pedidos = [
        (101, -6, "Mariana Souza", "(11) 90000-0001", "Bolo de pote", 6, 0, -4, "Pix", "Sim", "Entregue"),
        (102, -5, "Carlos Lima", "(11) 90000-0002", "Brigadeiro gourmet (cento)", 2, 20, -1, "Cartão de crédito", "Sim", "Entregue"),
        (102, -5, "Carlos Lima", "(11) 90000-0002", "Taxa de entrega", 1, 0, -1, "Cartão de crédito", "Sim", "Entregue"),
        (103, -3, "Ana Paula", "(11) 90000-0003", "Torta de limão inteira", 1, 0, -1, "Pix", "Não", "Pronto"),
        (104, -2, "Buffet Alegria", "(11) 90000-0004", "Kit festa 20 pessoas", 2, 30, 0, "Boleto", "Não", "Em produção"),
        (105, -1, "João Pedro", "(11) 90000-0005", "Bolo de pote", 10, 0, 2, "Dinheiro", "Não", "Novo"),
        (106, -1, "Fernanda Reis", "(11) 90000-0006", "Torta de limão inteira", 1, 0, 3, "Pix", "Sim", "Cancelado"),
    ]
    for i, (n, d, cli, tel, prod, q, desc, ent, pag, pago, st) in enumerate(pedidos):
        r = H + 1 + i
        p[f"A{r}"], p[f"B{r}"], p[f"C{r}"], p[f"D{r}"] = n, rel(d), cli, tel
        p[f"B{r}"].number_format = p[f"J{r}"].number_format = DATA
        p[f"E{r}"], p[f"F{r}"], p[f"H{r}"] = prod, q, desc
        p[f"J{r}"], p[f"K{r}"], p[f"L{r}"], p[f"M{r}"] = rel(ent), pag, pago, st

    # Painel com filtro de mês: B4 = mês, B5 = ano; período = [início, fim)
    ws = painel(wb, "Painel de Pedidos", [], linha_ini=4)
    ws["A4"], ws["B4"] = "Mês (1 a 12)", "=MONTH(TODAY())"
    ws["A5"], ws["B5"] = "Ano", "=YEAR(TODAY())"
    for ref in ("A4", "A5"):
        ws[ref].font = Font(bold=True)
    for ref in ("B4", "B5"):
        ws[ref].font = Font(size=13, bold=True, color="FFFFFF")
        ws[ref].fill = PatternFill("solid", fgColor=INDIGO)
        ws[ref].alignment = Alignment(horizontal="center")
    ws["C4"] = "← troque pelo mês que quer ver (a fórmula padrão mostra o mês atual)"
    ws["C4"].font = Font(italic=True, color=CINZA)
    ws.column_dimensions["C"].width = 60
    ini, fim = "DATE($B$5,$B$4,1)", "DATE($B$5,$B$4+1,1)"
    data, tot = "Pedidos!$B$5:$B$504", "Pedidos!$I$5:$I$504"
    st, pago = "Pedidos!$M$5:$M$504", "Pedidos!$L$5:$L$504"
    periodo = f'{data},">="&{ini},{data},"<"&{fim}'
    indicadores = [
        ("Faturamento do mês (sem cancelados)", f'=SUMIFS({tot},{periodo},{st},"<>Cancelado")', MOEDA),
        ("A receber no mês (Pago = Não)", f'=SUMIFS({tot},{periodo},{st},"<>Cancelado",{pago},"Não")', MOEDA),
        ("Itens vendidos no mês", f'=SUMIFS(Pedidos!$F$5:$F$504,{periodo},{st},"<>Cancelado")', "0"),
        ("Linhas canceladas no mês", f'=COUNTIFS({periodo},{st},"Cancelado")', "0"),
        ("Pedidos em aberto (todas as datas)",
         f'=SUMPRODUCT((Pedidos!$E$5:$E$504<>"")*({st}<>"Entregue")*({st}<>"Cancelado"))', "0"),
        ("Linhas ATRASADAS agora", '=COUNTIF(Pedidos!$N$5:$N$504,"ATRASADO")', "0"),
    ]
    for i, (rot, f, fmt) in enumerate(indicadores):
        r = 7 + i
        a, b = ws.cell(row=r, column=1, value=rot), ws.cell(row=r, column=2, value=f)
        a.font, b.font = Font(bold=True), Font(size=13, bold=True, color=INDIGO)
        b.number_format, b.alignment = fmt, Alignment(horizontal="right")
        a.border = b.border = BORDA
    r0 = 7 + len(indicadores) + 2
    ws.cell(row=r0, column=1, value="Vendas por produto no mês").font = Font(size=13, bold=True)
    for j, nome in enumerate(["Produto / serviço", "Quantidade", "Faturamento"], start=1):
        cel = ws.cell(row=r0 + 1, column=j, value=nome)
        cel.font, cel.fill = Font(bold=True, color=INDIGO), PatternFill("solid", fgColor=INDIGO_CLARO)
    ws.column_dimensions["C"].width = 60
    for i in range(30):
        r, src = r0 + 2 + i, H + 1 + i
        ws[f"A{r}"] = f"=IF('Catálogo'!A{src}=\"\",\"\",'Catálogo'!A{src})"
        ws[f"B{r}"] = f'=IF(A{r}="","",SUMIFS(Pedidos!$F$5:$F$504,Pedidos!$E$5:$E$504,A{r},{periodo},{st},"<>Cancelado"))'
        ws[f"C{r}"] = f'=IF(A{r}="","",SUMIFS({tot},Pedidos!$E$5:$E$504,A{r},{periodo},{st},"<>Cancelado"))'
        ws[f"C{r}"].number_format = MOEDA
        ws[f"C{r}"].alignment = Alignment(horizontal="left")
    return wb


# --------------------------------------------------------------------------
# Kit 3: Controle de vencimentos
# --------------------------------------------------------------------------
def kit_vencimentos():
    wb = Workbook()
    como_usar(wb, "Controle de Vencimentos", [
        "Na aba Vencimentos, liste tudo que tem data: impostos e guias, contas fixas, contratos, alvarás e licenças, "
        "certificado digital, seguros, mensalidades de sistemas.",
        "Preencha a Data de vencimento e, em Avisar com, quantos dias antes você quer ser alertado (padrão: 7).",
        "A Situação se atualiza todo dia sozinha: vermelho = VENCIDO, amarelo = VENCE EM BREVE, verde = no prazo.",
        "Quando pagar ou renovar, mude o Status para Feito. Se for recorrente, a coluna Próximo vencimento "
        "mostra a nova data: lance uma nova linha com ela.",
        "Abra a planilha toda segunda-feira e olhe o Painel: ele mostra o que vence nos próximos 7 e 30 dias e quanto vai sair do caixa.",
    ], [
        "As datas de impostos (DAS, DARF, guias estaduais e municipais) mudam conforme o regime e a cidade. "
        "Confirme as suas com o seu contador antes de cadastrar.",
        "Certificado digital e alvarás costumam levar dias para renovar: use um aviso maior (30 dias) para eles.",
        "Use a coluna Responsável quando mais de uma pessoa cuida das contas.",
    ])

    H = 4
    v = wb.create_sheet("Vencimentos")
    titulo(v, "Vencimentos", "Tudo que tem prazo: impostos, contas, contratos, licenças.")
    cols = [
        ("Item", 32, None, True), ("Categoria", 18, None, True), ("Responsável", 14, None, True),
        ("Data de vencimento", 13, DATA, True), ("Valor", 13, MOEDA, True), ("Recorrência", 13, None, True),
        ("Avisar com (dias)", 10, "0", True), ("Status", 11, None, True), ("Dias restantes", 10, "0", False),
        ("Situação", 17, None, False), ("Próximo vencimento", 13, DATA, False), ("Observação", 30, None, True),
    ]
    cabecalho(v, H, cols)
    formatos(v, H + 1, cols)
    lista(v, f"B{H + 1}:B{H + LINHAS}", ["Imposto / guia", "Conta fixa", "Contrato", "Alvará / licença",
                                         "Certificado digital", "Seguro", "Sistema / assinatura", "Fornecedor", "Outro"])
    lista(v, f"F{H + 1}:F{H + LINHAS}", ["Única", "Mensal", "Trimestral", "Semestral", "Anual"])
    lista(v, f"H{H + 1}:H{H + LINHAS}", ["Pendente", "Feito"])
    for r in range(H + 1, H + 1 + LINHAS):
        aviso = f'IF(G{r}="",7,G{r})'
        v[f"I{r}"] = f'=IF(OR(A{r}="",D{r}="",H{r}="Feito"),"",D{r}-TODAY())'
        v[f"J{r}"] = (f'=IF(OR(A{r}="",D{r}=""),"",IF(H{r}="Feito","FEITO",'
                      f'IF(I{r}<0,"VENCIDO",IF(I{r}<={aviso},"VENCE EM BREVE","No prazo"))))')
        v[f"K{r}"] = (f'=IF(OR(D{r}="",F{r}="",F{r}="Única"),"",EDATE(D{r},'
                      f'IF(F{r}="Mensal",1,IF(F{r}="Trimestral",3,IF(F{r}="Semestral",6,12)))))')
    faixa = f"J{H + 1}:J{H + LINHAS}"
    destaque(v, faixa, f'J{H + 1}="VENCIDO"', VERMELHO, VERMELHO_CLARO)
    destaque(v, faixa, f'J{H + 1}="VENCE EM BREVE"', AMARELO, AMARELO_CLARO)
    destaque(v, faixa, f'OR(J{H + 1}="No prazo",J{H + 1}="FEITO")', VERDE, VERDE_CLARO)
    exemplos = [
        ("DAS (Simples Nacional)", "Imposto / guia", "Contador", 5, 412.35, "Mensal", 5, "Pendente", "conferir valor com o contador"),
        ("Aluguel da loja", "Conta fixa", "Sócio", 2, 2800.0, "Mensal", 5, "Pendente", ""),
        ("Energia elétrica", "Conta fixa", "Sócio", -1, 389.9, "Mensal", 3, "Pendente", ""),
        ("Internet", "Conta fixa", "Sócio", 12, 129.9, "Mensal", 3, "Pendente", ""),
        ("Certificado digital A1", "Certificado digital", "Sócio", 25, 220.0, "Anual", 30, "Pendente", "renovar com antecedência"),
        ("Alvará de funcionamento", "Alvará / licença", "Contador", 60, 0, "Anual", 30, "Pendente", ""),
        ("Contrato do sistema de vendas", "Sistema / assinatura", "Sócio", 40, 99.0, "Mensal", 7, "Pendente", ""),
        ("Seguro do estabelecimento", "Seguro", "Sócio", -10, 1450.0, "Anual", 15, "Feito", "renovado"),
    ]
    for i, (item, cat, resp, dias, valor, rec, aviso, status, obs) in enumerate(exemplos):
        r = H + 1 + i
        v[f"A{r}"], v[f"B{r}"], v[f"C{r}"] = item, cat, resp
        v[f"D{r}"], v[f"E{r}"], v[f"F{r}"] = rel(dias), valor, rec
        v[f"G{r}"], v[f"H{r}"], v[f"L{r}"] = aviso, status, obs
        v[f"D{r}"].number_format = DATA

    d, val, st = "Vencimentos!$D$5:$D$504", "Vencimentos!$E$5:$E$504", "Vencimentos!$H$5:$H$504"
    sit = "Vencimentos!$J$5:$J$504"
    painel(wb, "Painel de Vencimentos", [
        ("Itens VENCIDOS", f'=COUNTIF({sit},"VENCIDO")', "0"),
        ("Itens que VENCEM EM BREVE", f'=COUNTIF({sit},"VENCE EM BREVE")', "0"),
        ("Valor vencido e não pago", f'=SUMIFS({val},{sit},"VENCIDO")', MOEDA),
        ("A pagar nos próximos 7 dias",
         f'=SUMIFS({val},{d},">="&TODAY(),{d},"<="&(TODAY()+7),{st},"<>Feito")', MOEDA),
        ("A pagar nos próximos 30 dias",
         f'=SUMIFS({val},{d},">="&TODAY(),{d},"<="&(TODAY()+30),{st},"<>Feito")', MOEDA),
        ("Itens pendentes cadastrados", f'=SUMPRODUCT((Vencimentos!$A$5:$A$504<>"")*({st}<>"Feito"))', "0"),
    ])
    return wb


KITS = {
    "controle-de-estoque": kit_estoque,
    "gestao-de-pedidos": kit_pedidos,
    "controle-de-vencimentos": kit_vencimentos,
}

if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for slug, fabrica in KITS.items():
        wb = fabrica()
        wb.active = 0
        destino = OUT / f"{slug}.xlsx"
        wb.properties.creator = "Resolvo Já"
        wb.properties.title = slug
        wb.save(destino)
        print(destino)
